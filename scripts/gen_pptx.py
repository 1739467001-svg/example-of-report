# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

ROOT="/home/user/example-of-report"
D=f"{ROOT}/assets/diagrams"; S=f"{ROOT}/assets/screenshots"
OUT=f"{ROOT}/slides/督导系统建设成果汇报.pptx"
NT=23  # total slides

NAVY=RGBColor(0x0E,0x2A,0x56); NAVY2=RGBColor(0x13,0x31,0x5F)
BLUE=RGBColor(0x25,0x63,0xEB); SKY=RGBColor(0x4A,0x90,0xE2)
TEAL=RGBColor(0x0E,0x9F,0x8E); GREEN=RGBColor(0x2B,0xA8,0x73)
GOLD=RGBColor(0xE8,0xA2,0x3D); CYAN=RGBColor(0x1F,0xA2,0xB8)
TEXT=RGBColor(0x1F,0x2A,0x44); GRAY=RGBColor(0x5C,0x6B,0x82)
LIGHT=RGBColor(0xEF,0xF3,0xFB); CARDBG=RGBColor(0xFF,0xFF,0xFF)
LINE=RGBColor(0xDC,0xE6,0xF4); CHIP=RGBColor(0xEA,0xF1,0xFF)
WHITE=RGBColor(0xFF,0xFF,0xFF); BLUEL=RGBColor(0x7F,0xB2,0xFF)
PHB=RGBColor(0xB9,0xCB,0xE6); PHT=RGBColor(0x7E,0x93,0xB5)
FONT="Microsoft YaHei"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=13.333,7.5; BLANK=prs.slide_layouts[6]
def slide(): return prs.slides.add_slide(BLANK)

def _font(run,size,color,bold,font):
    run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color; run.font.name=font
    rPr=run._r.get_or_add_rPr()
    for tag in ('a:ea','a:cs'):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set('typeface',font)

def text(sl,l,t,w,h,lines,size=18,color=TEXT,bold=False,align=PP_ALIGN.LEFT,
         font=FONT,anchor=MSO_ANCHOR.TOP,ls=1.0):
    tb=sl.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    if isinstance(lines,str): lines=[lines]
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=ls
        if isinstance(ln,list):
            for seg in ln:
                r=p.add_run(); r.text=seg[0]
                _font(r,seg[1] if len(seg)>1 else size,seg[2] if len(seg)>2 else color,
                      seg[3] if len(seg)>3 else bold,font)
        else:
            r=p.add_run(); r.text=ln; _font(r,size,color,bold,font)
    return tb

def rect(sl,l,t,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.08,shadow=False,dash=False):
    sp=sl.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else:
        sp.line.color.rgb=line; sp.line.width=Pt(lw)
        if dash:
            ln=sp.line._get_or_add_ln(); pd=ln.makeelement(qn('a:prstDash'),{'val':'dash'}); ln.append(pd)
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=radius
        except: pass
    sp.shadow.inherit=False
    if shadow:
        el=sp._element.spPr; ef=el.makeelement(qn('a:effectLst'),{}); el.append(ef)
        sh=ef.makeelement(qn('a:outerShdw'),{'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'})
        clr=sh.makeelement(qn('a:srgbClr'),{'val':'0E2A56'}); al=clr.makeelement(qn('a:alpha'),{'val':'16000'})
        clr.append(al); sh.append(clr); ef.append(sh)
    return sp

def bg(sl,color=LIGHT): rect(sl,0,0,SW,SH,fill=color,shape=MSO_SHAPE.RECTANGLE)
def fit_img(sl,path,bl,bt,bw,bh):
    iw,ih=Image.open(path).size; ar=iw/ih
    if ar>bw/bh: w=bw; h=bw/ar
    else: h=bh; w=bh*ar
    sl.shapes.add_picture(path,Inches(bl+(bw-w)/2),Inches(bt+(bh-h)/2),Inches(w),Inches(h))

def header(sl,tab,title,sub=""):
    rect(sl,0.62,0.42,0.12+0.135*len(tab),0.34,fill=BLUE,radius=0.5)
    text(sl,0.62,0.42,0.12+0.135*len(tab),0.34,tab,11,WHITE,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(sl,0.6,0.84,12.1,0.66,title,29,NAVY,True)
    rect(sl,0.64,1.52,0.86,0.055,fill=BLUE,radius=0.5)
    if sub: text(sl,0.64,1.63,12.0,0.4,sub,14,GRAY)

def footer(sl,n,dark=False):
    c=RGBColor(0x5F,0x7B,0xAB) if dark else RGBColor(0x9F,0xB0,0xC8)
    text(sl,0.6,7.08,8,0.3,"浙江工商大学研究生院督导系统 · 建设成果汇报",10,c)
    text(sl,11.4,7.08,1.3,0.3,f"{n} / {NT}",10,c,align=PP_ALIGN.RIGHT)

def bullets(sl,l,t,w,items,size=15,gap=0.62,dot=BLUE):
    y=t
    for it in items:
        rect(sl,l,y+0.07,0.12,0.12,fill=dot,shape=MSO_SHAPE.OVAL)
        text(sl,l+0.30,y-0.04,w-0.30,gap,it,size,TEXT,ls=1.15)
        y+=gap

def placeholder(sl,l,t,w,h,label):
    rect(sl,l,t,w,h,fill=RGBColor(0xF6,0xF9,0xFF),line=PHB,lw=1.6,radius=0.05,dash=True)
    text(sl,l,t+h/2-0.55,w,0.5,"［ 图 ］",30,PHB,True,PP_ALIGN.CENTER)
    text(sl,l,t+h/2+0.0,w,0.4,"建议替换为真实截图",17,RGBColor(0x5C,0x6B,0x82),True,PP_ALIGN.CENTER)
    text(sl,l,t+h/2+0.45,w,0.4,label,13,PHT,False,PP_ALIGN.CENTER)

def fdetail(n,tab,title,sub,items,right="ph",ph_label="",imgs=None):
    sl=slide(); bg(sl); header(sl,tab,title,sub)
    bullets(sl,0.7,2.55,5.0,items,size=16,gap=0.78)
    if right=="ph": placeholder(sl,6.1,2.5,6.6,3.9,ph_label)
    footer(sl,n); return sl

# 1 COVER
def s_cover():
    sl=slide(); sl.shapes.add_picture(f"{D}/cover-bg.png",0,0,Inches(SW),Inches(SH))
    rect(sl,1.0,1.7,0.62,0.62,fill=None,line=BLUEL,lw=1.4,radius=0.22)
    text(sl,1.0,1.7,0.62,0.62,"📖",22,BLUEL,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(sl,1.78,1.74,7,0.36,"浙江工商大学 · 研究生院",17,RGBColor(0xEA,0xF2,0xFF),True)
    text(sl,1.79,2.16,8,0.3,"ZHEJIANG GONGSHANG UNIVERSITY · GRADUATE SCHOOL",9.5,RGBColor(0x7E,0x9B,0xCE))
    text(sl,1.0,2.85,11,1.0,"研究生院督导系统",44,WHITE,True)
    text(sl,1.0,3.75,11,1.0,"建设成果汇报",44,BLUEL,True)
    rect(sl,1.05,4.8,1.0,0.08,fill=GOLD,radius=0.5)
    text(sl,1.0,5.0,11,0.5,"基于全校真实课表数据的研究生教学督导一体化平台",18,RGBColor(0xDB,0xE7,0xFB))
    text(sl,1.0,5.55,11,0.4,"数据驱动 · 流程闭环 · 多角色协同",13,RGBColor(0x8F,0xB0,0xE6))

# 2 AGENDA (7)
def s_agenda():
    sl=slide(); bg(sl); header(sl,"目录","汇报内容","CONTENTS")
    items=[("01","项目背景与目标","痛点与建设愿景"),("02","系统概览 · 数据底座","真实数据一比一还原"),
           ("03","四类角色 · 权限体系","权限层层递进"),("04","六大核心功能详解","逐项展开核心能力"),
           ("05","业务闭环与系统架构","看得见、管得住"),("06","成果 · 价值 · 规划","小结与展望"),
           ("07","答辩 Q&A","常见问题预演")]
    cw=5.85;ch=1.0;gx=0.4;gy=0.16;x0=0.62;y0=2.0
    for i,(n,t,d) in enumerate(items):
        r,c=divmod(i,2); x=x0+c*(cw+gx); y=y0+r*(ch+gy)
        rect(sl,x,y,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.10,shadow=True)
        text(sl,x+0.22,y+0.13,0.9,0.75,n,26,RGBColor(0xCD,0xD9,0xEE),True,anchor=MSO_ANCHOR.MIDDLE)
        rect(sl,x+1.15,y+0.27,0.46,0.46,fill=BLUE,radius=0.25)
        text(sl,x+1.15,y+0.25,0.46,0.46,"▣",15,WHITE,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,x+1.8,y+0.18,cw-1.95,0.45,t,16.5,NAVY,True)
        text(sl,x+1.8,y+0.58,cw-1.95,0.35,d,11.5,GRAY)
    footer(sl,2)

# 3 BACKGROUND
def s_background():
    sl=slide(); bg(sl); header(sl,"01 项目背景","为什么要建这个系统","传统研究生教学督导面临四大痛点")
    pains=[("找课难","课程分散在多学院、两校区，含单双周，难快速定位目标课"),
           ("评价散","依赖纸质评价表，填写、收集、归档环节多，结果难沉淀"),
           ("统计弱","评价数据散落纸面，学院与研究生院缺乏全局视图"),
           ("反馈慢","评价结果传递链路长，难成质量改进闭环")]
    cw=3.5;ch=1.55;gx=0.3;x0=0.62;y0=2.35
    for i,(t,d) in enumerate(pains):
        r,c=divmod(i,2); x=x0+c*(cw+gx); y=y0+r*(ch+0.28)
        rect(sl,x,y,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.08,shadow=True)
        rect(sl,x,y,0.10,ch,fill=BLUE,shape=MSO_SHAPE.RECTANGLE)
        text(sl,x+0.32,y+0.20,cw-0.5,0.45,t,18,BLUE,True)
        text(sl,x+0.32,y+0.68,cw-0.55,0.8,d,12.5,GRAY,ls=1.15)
    gx0=8.25; rect(sl,gx0,2.35,4.45,3.66,fill=NAVY,radius=0.06,shadow=True)
    text(sl,gx0+0.4,2.7,3.7,0.5,"建设目标",20,GOLD,True)
    text(sl,gx0+0.4,3.35,3.75,2.6,
         [[("打造一个 ",15,RGBColor(0xE8,0xF0,0xFF)),("数据驱动 · 流程闭环 · 多角色协同",15,WHITE,True),
          (" 的研究生教学督导一体化平台，把“全校总课表”变成 ",15,RGBColor(0xE8,0xF0,0xFF)),
          ("可检索、可督导、可评价、可统计",15,WHITE,True),(" 的在线平台。",15,RGBColor(0xE8,0xF0,0xFF))]],ls=1.5)
    footer(sl,3)

# 4 OVERVIEW
def s_overview():
    sl=slide(); bg(sl); header(sl,"02 系统概览","一个面向真实业务的督导平台",
                               "课程、教师、学院、校区、排课规则全部来源于学校真实数据")
    rect(sl,0.62,2.3,12.1,0.95,fill=CHIP,radius=0.10); rect(sl,0.62,2.3,0.10,0.95,fill=BLUE,shape=MSO_SHAPE.RECTANGLE)
    text(sl,0.95,2.42,11.6,0.7,[[("系统完全以学校《全校总课表》《角色信息表》《课程评价表》搭建，",15,NAVY2),
        ("所见即真实业务",15,BLUE,True),("，开箱即用。",15,NAVY2)]],anchor=MSO_ANCHOR.MIDDLE)
    fit_img(sl,f"{D}/data-foundation.png",0.62,3.55,12.1,3.2); footer(sl,4)

# generic image slide
def s_image(n,tab,title,sub,img):
    sl=slide(); bg(sl); header(sl,tab,title,sub)
    fit_img(sl,img,0.7,2.3,11.93,4.55); footer(sl,n)

# 6 FEATURES
def s_features():
    sl=slide(); bg(sl); header(sl,"04 核心功能","六大核心功能模块","覆盖督导工作全流程，下面逐项展开")
    feats=[("全校课程多维筛选","按周次/星期/学院/校区/教师组合筛选",BLUE),
           ("听课计划管理","待办 + 日历视图，课前自动提醒",SKY),
           ("标准化课程评价","20 项定量 + 定性，支持草稿",TEAL),
           ("数据看板与统计","关键指标总览，图形化展示",GREEN),
           ("消息通知与提醒","评价提交后自动通知",GOLD),
           ("工号登录与用户管理","工号唯一凭证，角色权限控制",CYAN)]
    cw=3.86;ch=1.95;gx=0.27;gy=0.3;x0=0.62;y0=2.35
    for i,(t,d,col) in enumerate(feats):
        r,c=divmod(i,3); x=x0+c*(cw+gx); y=y0+r*(ch+gy)
        rect(sl,x,y,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.08,shadow=True)
        rect(sl,x+0.3,y+0.28,0.62,0.62,fill=CHIP,radius=0.22)
        text(sl,x+0.3,y+0.26,0.62,0.62,"◆",18,col,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,x+0.3,y+1.02,cw-0.55,0.5,t,16,NAVY,True)
        text(sl,x+0.3,y+1.42,cw-0.55,0.5,d,12,GRAY,ls=1.12)
    footer(sl,6)

# 10 EVAL RECORD UI
def s_evalui():
    sl=slide(); bg(sl); header(sl,"04 核心功能 · 功能③","评价记录与界面","评价详情清晰呈现各维度得分、综合评分与星级")
    rect(sl,0.62,2.4,8.0,3.9,fill=CARDBG,line=LINE,lw=0.75,radius=0.04,shadow=True)
    fit_img(sl,f"{S}/eval-record.png",0.8,2.55,7.64,3.4)
    text(sl,0.62,6.05,8.0,0.35,"课程评价记录详情页（真实界面）",12,GRAY,align=PP_ALIGN.CENTER)
    rect(sl,8.95,2.4,3.75,3.9,fill=CARDBG,line=LINE,lw=0.75,radius=0.04,shadow=True)
    fit_img(sl,f"{S}/eval-form-item.png",9.15,3.4,3.35,1.7)
    text(sl,8.95,6.05,3.75,0.35,"标准化评价表项（分级评分）",12,GRAY,align=PP_ALIGN.CENTER)
    footer(sl,10)

# 11 DASHBOARD
def s_dashboard():
    sl=slide(); bg(sl); header(sl,"04 核心功能 · 功能④","研究生院主管工作台 · 数据可视化","关键指标一屏总览，评价结果多维度图形化呈现")
    rect(sl,0.62,2.4,7.7,3.95,fill=CARDBG,line=LINE,lw=0.75,radius=0.04,shadow=True)
    fit_img(sl,f"{S}/dashboard.png",0.8,2.55,7.34,3.5)
    text(sl,0.62,6.1,7.7,0.35,"研究生院主管数据看板（真实界面）",12,GRAY,align=PP_ALIGN.CENTER)
    rect(sl,8.6,2.4,4.1,3.95,fill=CARDBG,line=LINE,lw=0.75,radius=0.04,shadow=True)
    fit_img(sl,f"{D}/chart-eval-radar.png",8.75,2.55,3.8,3.5)
    text(sl,8.6,6.1,4.1,0.35,"课程评价多维度可视化（示例）",12,GRAY,align=PP_ALIGN.CENTER)
    footer(sl,11)

# 12 STATS
def s_stats():
    sl=slide(); bg(sl); header(sl,"04 核心功能 · 功能④","学院评价统计分析","按学院展示督导评价次数与评分分布，支撑质量决策")
    rect(sl,0.62,2.35,6.7,3.7,fill=CARDBG,line=LINE,lw=0.75,radius=0.04,shadow=True)
    fit_img(sl,f"{D}/chart-college.png",0.8,2.5,6.34,3.4)
    rect(sl,7.55,2.35,5.15,3.7,fill=CARDBG,line=LINE,lw=0.75,radius=0.04,shadow=True)
    fit_img(sl,f"{D}/chart-scoredist.png",7.7,2.5,4.85,3.4)
    text(sl,0.62,6.2,12.1,0.35,"※ 上图为示例数据，接入真实评价数据后由系统自动生成",12.5,PHT,align=PP_ALIGN.CENTER)
    footer(sl,12)

# 13 NOTIFY
def s_notify():
    sl=slide(); bg(sl); header(sl,"04 核心功能 · 功能⑤","消息通知与提醒","打通“评价完成 → 学院/研究生院知晓”的反馈链路")
    bullets(sl,0.7,2.6,5.0,["评价提交后自动通知研究生院主管与对应学院教学秘书",
        "听课计划课前自动站内提醒，避免漏听","通知可点击直达评价详情，反馈即时","让质量改进闭环真正跑起来"],size=16,gap=0.82)
    x=7.2;w=5.0
    def node(y,t,d,col,h=0.82):
        rect(sl,x,y,w,h,fill=CARDBG,line=LINE,lw=0.75,radius=0.12,shadow=True); rect(sl,x,y,0.1,h,fill=col,shape=MSO_SHAPE.RECTANGLE)
        text(sl,x+0.35,y+0.12,w-0.5,0.4,t,15,NAVY,True); text(sl,x+0.35,y+0.48,w-0.5,0.3,d,11,GRAY)
    node(2.5,"督导专家提交评价","完成课程评价",BLUE)
    text(sl,x,3.35,w,0.3,"▼",14,SKY,align=PP_ALIGN.CENTER)
    node(3.62,"系统自动推送通知","无需人工传递",SKY)
    text(sl,x,4.47,w,0.3,"▼",14,SKY,align=PP_ALIGN.CENTER)
    rect(sl,x,4.75,2.4,0.82,fill=CARDBG,line=LINE,lw=0.75,radius=0.12,shadow=True); rect(sl,x,4.75,0.1,0.82,fill=TEAL,shape=MSO_SHAPE.RECTANGLE)
    text(sl,x+0.3,4.87,2.1,0.4,"研究生院主管",13.5,NAVY,True); text(sl,x+0.3,5.22,2.1,0.3,"全局知晓",10.5,GRAY)
    rect(sl,x+2.6,4.75,2.4,0.82,fill=CARDBG,line=LINE,lw=0.75,radius=0.12,shadow=True); rect(sl,x+2.6,4.75,0.1,0.82,fill=GOLD,shape=MSO_SHAPE.RECTANGLE)
    text(sl,x+2.9,4.87,2.1,0.4,"对应学院教学秘书",12.5,NAVY,True); text(sl,x+2.9,5.22,2.1,0.3,"本院知晓",10.5,GRAY)
    footer(sl,13)

# 17 ACHIEVEMENTS
def s_ach():
    sl=slide(); bg(sl); header(sl,"06 建设成果","建设成果小结","六个维度，系统已具备真实可用能力")
    ach=[("真实数据驱动","1,344 门课 · 22 学院 · 2 校区 · 595 教师 · 58 用户"),
         ("角色全覆盖","四类角色权限清晰，数据隔离到位"),("流程全闭环","找课—听课—评价—统计—反馈一站打通"),
         ("评价标准化","复刻现行评价表，定量+定性，支持草稿与算分"),
         ("质量有保障","全部单元测试通过 · TypeScript 零类型错误"),("多端可用","电脑 / 平板 / 手机响应式适配")]
    cw=5.85;ch=1.18;gx=0.4;x0=0.62;y0=2.3
    for i,(t,d) in enumerate(ach):
        r,c=divmod(i,2); x=x0+c*(cw+gx); y=y0+r*(ch+0.22)
        rect(sl,x,y,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.10,shadow=True)
        rect(sl,x+0.25,y+0.30,0.58,0.58,fill=GREEN,radius=0.22)
        text(sl,x+0.25,y+0.28,0.58,0.58,"✓",20,WHITE,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,x+1.05,y+0.22,cw-1.2,0.5,t,17,NAVY,True); text(sl,x+1.05,y+0.66,cw-1.2,0.45,d,12,GRAY)
    footer(sl,17)

# 18 DEV
def s_dev():
    sl=slide(); bg(sl); header(sl,"06 开发方式","AI 辅助 + 数据驱动的敏捷开发","让“真实数据 → 可用系统”更快、更贴合一线需求")
    steps=[("1","真实数据自动建库","直接以学校 Excel 课表为输入，自动解析建立业务数据库，省去大量人工录入与结构设计"),
           ("2","快速搭建可用系统","四类角色、核心功能、商务蓝界面一次成型，前后端分离、PostgreSQL 稳定选型"),
           ("3","真实反馈快速迭代","围绕督导专家与研究生院真实使用反馈多轮打磨，持续修复优化，贴合实际")]
    cw=3.86;ch=3.0;gx=0.27;x0=0.62;y0=2.4
    for i,(n,t,d) in enumerate(steps):
        x=x0+i*(cw+gx); rect(sl,x,y0,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.07,shadow=True)
        rect(sl,x+0.35,y0+0.35,0.7,0.7,fill=BLUE,radius=0.5)
        text(sl,x+0.35,y0+0.33,0.7,0.7,n,24,WHITE,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,x+0.35,y0+1.25,cw-0.7,0.5,t,17,NAVY,True); text(sl,x+0.35,y0+1.85,cw-0.7,1.0,d,13,GRAY,ls=1.25)
    rect(sl,2.0,5.8,9.33,0.7,fill=RGBColor(0xFF,0xF7,0xEA),line=GOLD,lw=1.0,radius=0.3)
    text(sl,2.0,5.8,9.33,0.7,"显著压缩了从需求到可用系统的周期，使平台快速进入真实业务可用状态。",14,RGBColor(0xB5,0x70,0x10),True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    footer(sl,18)

# 19 VALUE
def s_value():
    sl=slide(); bg(sl); header(sl,"06 应用价值","为不同角色创造价值","一套系统，四方受益")
    val=[("督导专家","找课快、评价简、记录全；听课计划与提醒不漏听",BLUE),("学院","本院督导进度与评价质量一目了然",TEAL),
         ("研究生院","全局数据看板，督导工作有数可依、有据可查",SKY),("学校","研究生教学质量保障工作的数字化升级",NAVY)]
    cw=2.93;ch=3.3;gx=0.24;x0=0.62;y0=2.5
    for i,(t,d,col) in enumerate(val):
        x=x0+i*(cw+gx); rect(sl,x,y0,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.08,shadow=True)
        rect(sl,x,y0,cw,0.16,fill=col,shape=MSO_SHAPE.RECTANGLE)
        rect(sl,x+cw/2-0.45,y0+0.45,0.9,0.9,fill=CHIP,radius=0.5)
        text(sl,x+cw/2-0.45,y0+0.43,0.9,0.9,"●",26,col,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,x+0.2,y0+1.6,cw-0.4,0.5,t,18,NAVY,True,PP_ALIGN.CENTER)
        text(sl,x+0.25,y0+2.15,cw-0.5,1.0,d,12.5,GRAY,False,PP_ALIGN.CENTER,ls=1.3)
    footer(sl,19)

# 20 ROADMAP
def s_roadmap():
    sl=slide(); bg(sl); header(sl,"06 后续规划","下一步规划","让平台持续生长，覆盖更广、用得更深")
    road=[("角色扩展","增加分管院长等角色"),("课程扩展","导入 MBA 等专业学位课程及评价表"),
          ("报表导出","一键导出 PDF / Excel 评价报告，便于学期末汇总"),
          ("智能增强","探索语音评价自动转写打分，降低填写成本"),("分析升级","教师评价趋势、同学院对标分析等高级图表")]
    y0=2.4;rh=0.82;gap=0.12
    for i,(t,d) in enumerate(road):
        y=y0+i*(rh+gap); rect(sl,0.62,y,12.1,rh,fill=CARDBG,line=LINE,lw=0.75,radius=0.12,shadow=True)
        rect(sl,0.95,y+rh/2-0.1,0.2,0.2,fill=BLUE,radius=0.5,shape=MSO_SHAPE.OVAL)
        text(sl,1.4,y,2.6,rh,t,17,NAVY,True,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,4.0,y,8.5,rh,d,14,GRAY,anchor=MSO_ANCHOR.MIDDLE)
    footer(sl,20)

# 21/22 QA
def s_qa(n,tab,title,sub,items):
    sl=slide(); bg(sl); header(sl,tab,title,sub)
    cw=5.85;ch=2.0;gx=0.4;gy=0.22;x0=0.62;y0=2.2
    for i,(q,a) in enumerate(items):
        r,c=divmod(i,2); x=x0+c*(cw+gx); y=y0+r*(ch+gy)
        rect(sl,x,y,cw,ch,fill=CARDBG,line=LINE,lw=0.75,radius=0.07,shadow=True)
        rect(sl,x+0.28,y+0.26,0.42,0.42,fill=BLUE,radius=0.22)
        text(sl,x+0.28,y+0.24,0.42,0.42,"Q",15,WHITE,True,PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(sl,x+0.85,y+0.24,cw-1.1,0.7,q,15,NAVY,True,ls=1.1)
        text(sl,x+0.32,y+1.0,cw-0.6,1.0,a,12,GRAY,ls=1.25)
    footer(sl,n)

# 23 CLOSING
def s_closing():
    sl=slide(); sl.shapes.add_picture(f"{D}/cover-bg.png",0,0,Inches(SW),Inches(SH))
    text(sl,0,2.5,SW,1.4,"谢　谢",54,WHITE,True,PP_ALIGN.CENTER)
    rect(sl,SW/2-0.6,4.05,1.2,0.08,fill=GOLD,radius=0.5)
    text(sl,0,4.4,SW,0.6,"浙江工商大学研究生院督导系统",22,RGBColor(0xDB,0xE7,0xFB),False,PP_ALIGN.CENTER)
    text(sl,0,5.1,SW,0.4,"数据驱动 · 流程闭环 · 多角色协同",14,RGBColor(0x8F,0xB0,0xE6),False,PP_ALIGN.CENTER)

# ---- build in order ----
s_cover(); s_agenda(); s_background(); s_overview()
s_image(5,"03 角色与权限","四类角色 · 权限层层递进","数据隔离 · 只读保护 · 逐级汇总，贴合研究生院真实管理结构",f"{D}/role-matrix.png")
s_features()
fdetail(7,"04 核心功能 · 功能①","全校课程多维筛选","按多维度精确定位目标课程，与原始课表完全一致",
        ["全校 1,344 门课程在线检索，告别翻纸质课表","支持 周次 / 星期 / 学院 / 校区 / 教师 多维组合筛选",
         "筛选逻辑与原课表一致（如“会计学院 + 第 3 周”精确联动）","督导专家可快速锁定“本周哪天、哪个学院、哪位老师”有课"],
        ph_label="全校课程 · 筛选页面")
fdetail(8,"04 核心功能 · 功能②","听课计划 · 待办与日历视图","把有意向的课程一键纳入计划，按周直观安排",
        ["一键将课程加入听课计划（类似待办清单）","日历视图：以“节次 × 星期”网格展示当周安排",
         "待听课 / 已评价 / 已取消 三种状态不同颜色区分","课前自动站内提醒，避免漏听"],
        ph_label="听课计划 · 日历视图")
s_image(9,"04 核心功能 · 功能③","标准化课程评价体系","完全复刻学校现行《课程评价表》：20 项定量 + 定性，系统自动计算综合评分",f"{D}/eval-system.png")
s_evalui(); s_dashboard(); s_stats(); s_notify()
fdetail(14,"04 核心功能 · 功能⑥","工号登录与用户管理","以工号为唯一凭证，角色权限清晰可控",
        ["全体用户以工号作为唯一登录凭证","初始密码即工号，支持首次登录后自助改密",
         "基于角色的权限控制与数据隔离（如学院秘书仅见本院）","研究生院主管可进行用户与角色管理"],
        ph_label="工号登录页面")
s_image(15,"05 业务闭环","督导业务闭环","从找课到改进，串成一条可沉淀、可统计、可追溯的数字化链路",f"{D}/business-loop.png")
s_image(16,"05 技术架构","系统架构（简要）","前后端分离 · 数据驱动 · 真实课表自动解析建库",f"{D}/architecture.png")
s_ach(); s_dev(); s_value(); s_roadmap()
s_qa(21,"07 答辩 Q&A","常见问题预演（一）","建设与数据",[
  ("数据来源是否可靠？如何保证与真实课表一致？","全部数据来源于学校《全校总课表》等真实文件，自动解析建库（1,344 门课、22 学院、595 名教师），筛选逻辑与原课表一致，可随新学期课表定期更新。"),
  ("评价数据安全吗？会不会被篡改？","工号唯一登录 + 角色权限控制；学院秘书对评价为只读、不可修改删除；数据存于 PostgreSQL，操作可追溯。"),
  ("评价标准如何统一？","评价表完全复刻学校现行《课程评价表》，20 项定量 + 定性，系统自动计算综合评分，口径一致、结果可比。"),
  ("能否与现有信息化系统对接？","前后端分离、以工号为唯一标识，预留数据导入与接口能力，可逐步与研究生管理等系统对接。"),
])
s_qa(22,"07 答辩 Q&A","常见问题预演（二）","应用与推广",[
  ("老师上手难不难？","工号即账号、初始密码为工号；多维筛选快速找课，日历视图 + 课前提醒，评价支持保存草稿分次完成，上手成本低。"),
  ("研究生院如何掌握全校督导情况？","主管数据看板一屏总览（课程数、已完成评价、督导专家数、覆盖学院），并按学院图形化展示评价次数与评分分布。"),
  ("能否支持 MBA 等其他课程？","可扩展。后续计划导入 MBA 等专业学位课程及对应评价表。"),
  ("下一步规划是什么？","分管院长角色、PDF/Excel 报表导出、语音评价自动打分、教师评价趋势与同学院对标分析。"),
])
s_closing()

prs.save(OUT)
print("saved PPTX:",OUT,"slides:",len(prs.slides._sldIdLst))
